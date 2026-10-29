# myapp/Card/docx_generator.py
import os
import sys


def generate_cards_docx(card_pairs, output, back_first=True, card_width_mm=51.0, card_height_mm=86.8):
    """
    Generates a Microsoft Word document (.docx) with cards arranged in a 3-column Table Grid,
    matching NIELIT standard card export layout:
      Row 0: [Card 1 Back, Card 1 Front, Card 2 Back]
      Row 1: [Card 2 Front, Card 3 Back, Card 3 Front]
      ...

    card_pairs: list of dicts with {'front': path, 'back': path, 'title': str}
    output: file path or BytesIO buffer
    back_first: if True, places Back then Front (matching exact template), else Front then Back.
    """
    try:
        import docx
        from docx.shared import Mm, Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT
    except ImportError as e:
        raise ImportError(
            f"python-docx is not loaded in this worker ({e}). "
            f"Please go to PythonAnywhere 'Web' tab and click the green 'Reload' button."
        )

    doc = docx.Document()

    # Page setup - A4 with standard margins
    sec = doc.sections[0]
    sec.page_width = Mm(210)
    sec.page_height = Mm(297)
    sec.top_margin = Mm(12.7)
    sec.bottom_margin = Mm(12.7)
    sec.left_margin = Mm(14.0)
    sec.right_margin = Mm(14.0)

    # Collect flat list of images in order
    image_list = []
    for pair in card_pairs:
        front = pair.get('front')
        back = pair.get('back')
        if back_first:
            if back and os.path.exists(back):
                image_list.append(back)
            if front and os.path.exists(front):
                image_list.append(front)
        else:
            if front and os.path.exists(front):
                image_list.append(front)
            if back and os.path.exists(back):
                image_list.append(back)

    if not image_list:
        return False

    # Create 3-column table
    table = doc.add_table(rows=0, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    table.autofit = False

    col_w = Mm(58.0)
    for col in table.columns:
        col.width = col_w

    # Populate rows (3 images per row)
    for i in range(0, len(image_list), 3):
        row = table.add_row()
        for j in range(3):
            cell = row.cells[j]
            cell.width = col_w
            img_idx = i + j
            if img_idx < len(image_list):
                img_path = image_list[img_idx]
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                run = p.add_run()
                run.add_picture(img_path, width=Mm(card_width_mm), height=Mm(card_height_mm))

    doc.save(output)
    return True
